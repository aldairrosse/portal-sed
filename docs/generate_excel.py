"""Genera Excel con las tablas de Solicitud de Requisitos y WBS."""

import re
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

BASE = Path(__file__).resolve().parent.parent
OUTPUT = BASE / "docs" / "sed-requisitos-wbs.xlsx"

# --- Estilos ---
HEADER_FONT = Font(name="Calibri", bold=True, size=11, color="FFFFFF")
HEADER_FILL = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
SECTION_FONT = Font(name="Calibri", bold=True, size=11, color="2F5496")
SECTION_FILL = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")
NORMAL_FONT = Font(name="Calibri", size=10)
WRAP = Alignment(wrap_text=True, vertical="top")
THIN_BORDER = Border(
    left=Side(style="thin", color="B4C6E7"),
    right=Side(style="thin", color="B4C6E7"),
    top=Side(style="thin", color="B4C6E7"),
    bottom=Side(style="thin", color="B4C6E7"),
)


def parse_md_table(text: str) -> list[list[str]]:
    """Extrae filas de una tabla markdown, excluyendo separadores ---."""
    rows = []
    for line in text.strip().splitlines():
        line = line.strip()
        if not line.startswith("|"):
            continue
        cells = [c.strip() for c in line.split("|")]
        # split agrega strings vacíos al inicio/final
        cells = cells[1:-1] if len(cells) > 2 else cells
        # saltar separador
        if all(re.match(r"^[-:]+$", c) for c in cells):
            continue
        rows.append(cells)
    return rows


def extract_section_tables(md_text: str) -> dict[str, list[list[str]]]:
    """Extrae todas las tablas de secciones del markdown."""
    tables = {}
    current_section = None
    current_table_lines = []

    def flush():
        nonlocal current_table_lines
        if current_section and current_table_lines:
            tbl = parse_md_table("\n".join(current_table_lines))
            if tbl:
                tables[current_section] = tbl
        current_table_lines = []

    for line in md_text.splitlines():
        # Detectar encabezados de sección (## Sección X: Nombre)
        m = re.match(r"^##\s+Sección\s+\d+:\s+(.+)", line)
        if m:
            flush()
            current_section = m.group(1).strip()
            continue
        # Detectar tablas
        if line.strip().startswith("|"):
            current_table_lines.append(line)
        else:
            if current_table_lines:
                flush()
                current_table_lines = []

    flush()
    return tables


def extract_all_tables(md_text: str) -> dict[str, list[list[str]]]:
    """Extrae todas las tablas del markdown usando encabezados ## como clave."""
    tables = {}
    current_section = None
    current_table_lines = []

    def flush():
        nonlocal current_table_lines
        if current_section and current_table_lines:
            tbl = parse_md_table("\n".join(current_table_lines))
            if tbl:
                tables[current_section] = tbl
        current_table_lines = []

    for line in md_text.splitlines():
        # Detectar encabezados de sección (## Nombre)
        m = re.match(r"^##\s+(.+)", line)
        if m:
            flush()
            current_section = m.group(1).strip()
            continue
        # Detectar tablas
        if line.strip().startswith("|"):
            current_table_lines.append(line)
        else:
            if current_table_lines:
                flush()
                current_table_lines = []

    flush()
    return tables


def style_header(ws, row_num: int, ncols: int):
    for col in range(1, ncols + 1):
        cell = ws.cell(row=row_num, column=col)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = WRAP
        cell.border = THIN_BORDER


def style_section_row(ws, row_num: int, ncols: int):
    for col in range(1, ncols + 1):
        cell = ws.cell(row=row_num, column=col)
        cell.font = SECTION_FONT
        cell.fill = SECTION_FILL
        cell.alignment = WRAP
        cell.border = THIN_BORDER


def style_data_row(ws, row_num: int, ncols: int):
    for col in range(1, ncols + 1):
        cell = ws.cell(row=row_num, column=col)
        cell.font = NORMAL_FONT
        cell.alignment = WRAP
        cell.border = THIN_BORDER


def write_solicitud(wb: Workbook, md_text: str):
    """Escribe una tabla limpia unificada de requisitos."""
    ws = wb.active
    ws.title = "Solicitud de Requisitos"

    tables = extract_section_tables(md_text)

    row = 1
    # Título
    ws.cell(row=row, column=1, value="SED — Solicitud de Requisitos").font = Font(
        name="Calibri", bold=True, size=14, color="2F5496"
    )
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
    row += 2

    # Encabezados de la tabla unificada: #, Sección, Quién, Historia, Cómo, Para qué, Criterios de aceptación
    headers = ["#", "Sección", "Quién", "Historia", "Cómo", "Para qué", "Criterios de aceptación"]
    for i, hdr in enumerate(headers, 1):
        ws.cell(row=row, column=i, value=hdr)
    style_header(ws, row, len(headers))
    row += 1

    # Recorrer cada sección y sus requisitos (excluir infraestructura)
    skip_sections = {"Infraestructura del Sistema"}
    for section_name, tbl in tables.items():
        if section_name in skip_sections:
            continue
        # La primera fila es el encabezado original (#, Quién, Historia, Cómo, Para qué, Criterios)
        # Las siguientes filas son los datos
        for data_row in tbl[1:]:
            # data_row[0] = #, data_row[1] = Quién, data_row[2] = Historia, etc.
            ws.cell(row=row, column=1, value=data_row[0] if len(data_row) > 0 else "")
            ws.cell(row=row, column=2, value=section_name)
            ws.cell(row=row, column=3, value=data_row[1] if len(data_row) > 1 else "")
            ws.cell(row=row, column=4, value=data_row[2] if len(data_row) > 2 else "")
            ws.cell(row=row, column=5, value=data_row[3] if len(data_row) > 3 else "")
            ws.cell(row=row, column=6, value=data_row[4] if len(data_row) > 4 else "")
            ws.cell(row=row, column=7, value=data_row[5] if len(data_row) > 5 else "")
            style_data_row(ws, row, len(headers))
            row += 1

    # Ajustar anchos
    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["C"].width = 20
    ws.column_dimensions["D"].width = 45
    ws.column_dimensions["E"].width = 40
    ws.column_dimensions["F"].width = 35
    ws.column_dimensions["G"].width = 55


def write_wbs(wb: Workbook, md_text: str):
    """Escribe una tabla limpia unificada del WBS."""
    ws = wb.create_sheet("WBS")

    tables = extract_all_tables(md_text)

    row = 1
    # Título
    ws.cell(row=row, column=1, value="SED — WBS del Proyecto").font = Font(
        name="Calibri", bold=True, size=14, color="2F5496"
    )
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
    row += 2

    # Tabla unificada de tareas
    ws.cell(row=row, column=1, value="Detalle de Tareas").font = SECTION_FONT
    style_section_row(ws, row, 6)
    row += 1

    headers = ["ID", "Sección", "Tarea", "Dependencias", "Horas", "Comentarios"]
    for i, hdr in enumerate(headers, 1):
        ws.cell(row=row, column=i, value=hdr)
    style_header(ws, row, len(headers))
    row += 1

    # Recorrer secciones 1-11
    for section_num in range(1, 12):
        for key, tbl in tables.items():
            if key.startswith(f"{section_num}.") or key.startswith(f"{section_num} ") or key.startswith(f"{section_num}. "):
                # Saltar Resumen Ejecutivo y otros no-numéricos
                if key == "Resumen Ejecutivo" or key == "Archivos de Referencia" or key == "Diagrama de Dependencias":
                    continue
                # Extraer nombre limpio de la sección
                section_label = re.sub(r"^\d+\.\s*", "", key).strip()
                # Las filas de datos (sin header)
                for data_row in tbl[1:]:
                    ws.cell(row=row, column=1, value=data_row[0] if len(data_row) > 0 else "")
                    ws.cell(row=row, column=2, value=section_label)
                    ws.cell(row=row, column=3, value=data_row[1] if len(data_row) > 1 else "")
                    ws.cell(row=row, column=4, value=data_row[2] if len(data_row) > 2 else "")
                    ws.cell(row=row, column=5, value=data_row[3] if len(data_row) > 3 else "")
                    ws.cell(row=row, column=6, value=data_row[4] if len(data_row) > 4 else "")
                    style_data_row(ws, row, len(headers))
                    row += 1

    # Ajustar anchos
    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["C"].width = 45
    ws.column_dimensions["D"].width = 18
    ws.column_dimensions["E"].width = 10
    ws.column_dimensions["F"].width = 60


def main():
    solicitud_md = (BASE / "docs" / "solicitud-requisitos.md").read_text(encoding="utf-8")
    wbs_md = (BASE / "docs" / "wbs-proyecto.md").read_text(encoding="utf-8")

    wb = Workbook()
    write_solicitud(wb, solicitud_md)
    write_wbs(wb, wbs_md)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT)
    print(f"Excel generado: {OUTPUT}")


if __name__ == "__main__":
    main()
